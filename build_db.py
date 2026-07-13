#!/usr/bin/env python3
"""학원 관리 프로토타입 DB(엑셀) 생성 — v3.

앱을 구동하는 하드코딩(src/data.ts)의 전체 내용을 DB에 그대로 적재.
(원생 + 예비원생 296명, 상담 이력 전체) + 8월 청구로 monthly_fee 병합.

시트:
  students       : 학생 마스터(원생+예비원생). 앱의 모든 필드 포함
  consultations  : 상담 이력 (notes → 행 단위)
  schools        : 학교 lookup (지정값, 6월→8월 / 12월→1월)
  classes        : 수업 마스터 (지정값)
  teachers       : 강사 lookup (T01~T04, 이름 미확인)

출력: ~/Desktop/학원관리_DB.xlsx
"""
import csv
import json
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

ROOT = os.path.dirname(os.path.abspath(__file__))
OUT = os.path.expanduser("~/Desktop/학원관리_DB.xlsx")
ENROLLED_STATUS = "원생 (등록완료)"

# ---------------------------------------------------------------- 지정 데이터
YEAR_END_MAP = {"6월": "8월", "12월": "1월", "": ""}
SCHOOLS_RAW = [
    ("SCH001", "ACSI Int", "", "12월"), ("SCH002", "SJI Int", "", "12월"),
    ("SCH003", "SJI Independent", "", "12월"), ("SCH004", "ACSI Independent", "", "12월"),
    ("SCH005", "OFS", "", "6월"), ("SCH006", "UWC East", "", "6월"),
    ("SCH007", "UWC Dover", "", "6월"), ("SCH008", "SAS", "", "6월"),
    ("SCH009", "AIS", "", ""), ("SCH010", "CIS", "", "6월"),
    ("SCH011", "Chatsworth", "", "6월"), ("SCH012", "DCIS", "", "6월"),
    ("SCH013", "Dulwich", "", ""), ("SCH014", "Etonhouse", "", "6월"),
    ("SCH015", "GESS", "", "6월"), ("SCH016", "GIIS", "", ""),
    ("SCH017", "HCIS", "", "12월"), ("SCH018", "ICS", "", "6월"),
    ("SCH019", "ISS", "", "6월"), ("SCH020", "Nexus", "", "6월"),
    ("SCH021", "One World Int", "", ""), ("SCH022", "SOTA", "", "12월"),
    ("SCH023", "SKIS", "", ""), ("SCH024", "SAIS", "", "6월"),
    ("SCH025", "TTS", "", "6월"), ("SCH026", "NLCS", "", "6월"),
]
SCHOOLS = [(sid, name, note, YEAR_END_MAP.get(ye, ye)) for sid, name, note, ye in SCHOOLS_RAW]

CLASSES = [
    ("c0001", "G11 IB AA SL Nov", "IB", "Math AA SL", "T02", 130),
    ("c0002", "G10 Add Math Nov", "IGCSE", "Math", "T02", 110),
    ("C001", "G9 Add Math Advanced Nov", "IGCSE", "Math", "T02", 110),
    ("C002", "G11 IB AA SL", "IB", "Math", "T02", 130),
    ("C003", "G12 IB AA SL", "IB", "Math", "T02", 130),
    ("C004", "G11 IB AA HL A반", "IB", "Math", "T02", 130),
    ("C005", "G9 Add Math A반", "IGCSE", "Math", "T02", 110),
    ("C006", "G10 Add Math A반", "IGCSE", "Math", "T02", 110),
    ("C007", "G9 Add Math A반", "IGCSE", "Add Math", "T02", 110),
    ("C008", "G11 IB AA SL 통합반", "IB", "Math AA SL", "T02", 130),
    ("C009", "G8 Math A반", "Local Math", "Math", "T02", 100),
    ("C010", "G12 IB AA SL 통합반", "IB", "Math", "T02", 130),
    ("C011", "G11 IB AA HL OFS", "IB", "Math AA HL", "T02", 130),
    ("C012", "G11 IB AA HL UWC East", "IB", "Math AA HL", "T02", 130),
    ("C013", "G10 Add Math Nov", "IGCSE", "Math", "T02", 110),
    ("C014", "G12 IB AA HL Nov", "IB", "Math", "T01", 130),
    ("C015", "G12 IB AA HL OFS", "IB", "Math", "T01", 130),
    ("C016", "G12 IB AA HL UWC Dover 1", "IB", "Math", "T01", 130),
    ("C017", "G12 IB AA HL UWC Dover 2", "IB", "Math", "T01", 130),
    ("C018", "G11 IB AA HL Nov A반", "IB", "Math", "T01", 130),
    ("C019", "G11 IB AA HL Nov B반", "IB", "Math", "T01", 130),
    ("C020", "G11 IB AA HL UWC Dover 1", "IB", "Math", "T01", 130),
    ("C021", "G11 IB AA HL UWC Dover 2", "IB", "Math", "T01", 130),
    ("C022", "G10 Add Math B반", "IGCSE", "Math", "T01", 110),
    ("C023", "G10 Add Math C반", "IGCSE", "Math", "T01", 130),
    ("C024", "G10 IB AA HL 선행반", "IB", "Math", "T01", 110),
    ("C025", "G9 Add Math C반", "IGCSE", "Math", "T01", 110),
    ("C026", "G9 Ext Math", "IGCSE", "Ext Math", "T01", 110),
    ("C027", "G3 Math P3 Local Math", "Local Math", "Math", "T01", None),
    ("C028", "G9 Add Math B반", "IGCSE", "Math", "T03", 110),
    ("C029", "G12 IB AA HL UWC Dover", "IB", "Math", "T01", 114),
    ("C030", "G12 IB AA HL OFS", "IB", "Math", "T01", 114),
    ("C031", "G12 IB AA HL", "IB", "Math", "T01", 114),
    ("C032", "G11 IB AA HL B반", "IB", "Math", "T01", 114),
    ("C033", "G8 Math", "기타", "Math", "T03", 95),
    ("C034", "G5 Math Pre-Algebra", "기타", "Math", "T03", 95),
    ("C035", "G10 Add Math S반", "IGCSE", "Math", "T02", 110),
    ("C036", "G10 Add Math A반", "IGCSE", "Math", "T02", 110),
    ("C037", "G7 Math A반", "기타", "Math", "T04", None),
    ("C038", "G7 Math B반", "기타", "Math", "T04", None),
    ("C039", "G7 Math C반", "기타", "Math", "T04", None),
    ("C040", "G8 Math B반", "Local Math", "Math", "T01", None),
    ("C041", "G8 Math B반", "Local Math", "Math", "T03", None),
    ("C042", "G7 Math Nov", "Local Math", "Math", "T03", None),
    ("C043", "G5 Math Pre-Algebra A반", "기타", "Math", "T03", None),
    ("C044", "G5 Math Pre-Algebra B반", "기타", "Math", "T03", None),
]
TEACHERS = [
    ("T01", "", "미확인 - 실제 강사명 매핑 필요"),
    ("T02", "", "미확인 - 실제 강사명 매핑 필요"),
    ("T03", "", "미확인 - 실제 강사명 매핑 필요"),
    ("T04", "", "미확인 - 실제 강사명 매핑 필요"),
]


# ---------------------------------------------------------------- helpers
def load_data_ts():
    txt = open(os.path.join(ROOT, "src", "data.ts"), encoding="utf-8").read()
    return json.loads(txt[txt.index("= [") + 2: txt.rindex("]") + 1])


def to_num(s):
    s = (s or "").replace(",", "").replace("$", "").strip()
    if s in ("", "-"):
        return None
    try:
        return int(float(s))
    except ValueError:
        return None


def parse_aug_totals():
    """8월 청구 → 학생명별 total 합계 (monthly_fee 병합용)."""
    rows = list(csv.reader(open(os.path.join(ROOT, "aug.csv"), encoding="utf-8")))
    hi = next(i for i, r in enumerate(rows) if r and r[0].strip() == "no.")
    totals = {}
    for r in rows[hi + 1:]:
        if len(r) < 14 or not r[3].strip():
            continue
        totals[r[3].strip()] = totals.get(r[3].strip(), 0) + (to_num(r[13]) or 0)
    return totals


def split_name(name):
    base = re.split(r"[(（]", name or "", 1)[0].strip()
    if re.search(r"[가-힣]", base):
        return "", base
    return base, ""


def program_of(cls):
    c = cls or ""
    if re.search(r"IB|AA HL|AA SL", c):
        return "IB"
    if re.search(r"Add Math|Ext Math|IGCSE", c):
        return "IGCSE"
    if "Local" in c:
        return "Local Math"
    return "기타" if c and c != "-" else ""


def subject_of(cls):
    c = cls or ""
    if "AA HL" in c:
        return "Math AA HL"
    if "AA SL" in c:
        return "Math AA SL"
    if "Add Math" in c:
        return "Add Math"
    if "Ext Math" in c:
        return "Ext Math"
    return "Math" if c and c != "-" else ""


def grade_num(s):
    m = re.search(r"G\s*(\d+)", s or "", re.I)
    return int(m.group(1)) if m else 10**9


def snum(sn):
    m = re.search(r"(\d+)\s*$", sn or "")
    return int(m.group(1)) if m else 10**9


# ---------------------------------------------------------------- 데이터 적재
data = load_data_ts()
aug_totals = parse_aug_totals()

# student_id 부여: 원생=학생번호, 예비원생=INQ-####
def student_id_of(s):
    return s["studentNumber"] if s.get("studentNumber") else f"INQ-{int(s['id']):04d}"

# 정렬: 원생(학생번호 오름차순) → 예비원생(문의일 내림차순)
enrolled = sorted([s for s in data if s["status"] == ENROLLED_STATUS], key=lambda s: snum(s.get("studentNumber")))
prospective = sorted([s for s in data if s["status"] != ENROLLED_STATUS], key=lambda s: s.get("date", ""), reverse=True)
ordered = enrolled + prospective

student_rows = []
consult_rows = []
cid = 0
for s in ordered:
    sid = student_id_of(s)
    en, kr = split_name(s["name"])
    cls = s.get("className", "")
    student_rows.append([
        sid, en, kr,
        s.get("school", ""), s.get("grade", ""),
        program_of(cls), subject_of(cls),
        (cls if cls and cls != "-" else ""),   # class_name (반)
        s["status"],
        s.get("date", ""),                      # first_inquiry_date
        s.get("studentPhone", ""),              # student_contact
        s.get("parentPhone", ""),               # parent_contact
        "",                                     # parent_name (원본 없음)
        s.get("email", ""),
        s.get("memo", ""),
        aug_totals.get(s["name"]),              # monthly_fee (8월 청구)
    ])
    for n in s.get("notes", []):
        cid += 1
        consult_rows.append([
            cid, sid, s["name"],
            n.get("date", ""), n.get("method", ""), n.get("content", ""),
        ])


# ---------------------------------------------------------------- xlsx write
HEADER_FILL = PatternFill("solid", fgColor="1F2937")
HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
THIN = Side(style="thin", color="E5E7EB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
wb = Workbook()


def add_sheet(title, headers, rows, widths, table_name, number_cols=None):
    ws = wb.create_sheet(title)
    ws.append(headers)
    for cell in ws[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center")
    for r in rows:
        ws.append(r)
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=len(headers)):
        for c in row:
            c.border = BORDER
            c.alignment = Alignment(vertical="center")
    for col in (number_cols or []):
        for r in range(2, ws.max_row + 1):
            ws.cell(row=r, column=col).number_format = "#,##0"
    last = f"{get_column_letter(len(headers))}{max(ws.max_row, 2)}"
    tbl = Table(displayName=table_name, ref=f"A1:{last}")
    tbl.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    ws.add_table(tbl)
    ws.freeze_panes = "A2"
    return ws


# README
ws = wb.active
ws.title = "README"
for r in [
    ["학원 관리 프로토타입 DB (v3 · 앱 하드코딩 전체 적재)", ""],
    ["", ""],
    ["시트(테이블)", "설명 / 키"],
    ["students", "학생 마스터(원생+예비원생 전체). PK=student_id (원생=학생번호, 예비원생=INQ-####)"],
    ["consultations", "상담 이력. FK=student_id→students. notes를 행 단위로 전개"],
    ["schools", "학교 lookup. PK=school_id. year_end: 6월→8월, 12월→1월"],
    ["classes", "수업 마스터(지정값). PK=class_id, teacher_id→teachers"],
    ["teachers", "강사 lookup. PK=teacher_id(T01~T04). 이름 미확인 — 매핑 필요"],
    ["", ""],
    ["비고", "status는 7단계 그대로. program/subject/class_name은 수강 반 기준."],
    ["", "monthly_fee는 8월 청구 합계(매칭된 경우). student_contact=학생, parent_contact=학부모 번호."],
]:
    ws.append(r)
ws["A1"].font = Font(bold=True, size=15, color="1F2937")
ws["A3"].font = ws["B3"].font = Font(bold=True, color="FFFFFF")
ws["A3"].fill = ws["B3"].fill = HEADER_FILL
ws.column_dimensions["A"].width = 16
ws.column_dimensions["B"].width = 96
for row in ws.iter_rows(min_row=4, max_row=ws.max_row, max_col=2):
    for c in row:
        c.alignment = Alignment(vertical="center", wrap_text=True)

add_sheet("students",
          ["student_id", "name_en", "name_kr", "school", "grade", "program", "subject",
           "class_name", "status", "first_inquiry_date", "student_contact",
           "parent_contact", "parent_name", "email", "memo", "monthly_fee"],
          student_rows,
          [16, 15, 12, 13, 7, 11, 12, 24, 24, 15, 16, 16, 12, 28, 24, 12],
          "tbl_students", number_cols=[16])

add_sheet("consultations",
          ["consultation_id", "student_id", "student_name", "date", "method", "content"],
          consult_rows,
          [15, 16, 14, 13, 16, 90],
          "tbl_consultations", number_cols=[1])

add_sheet("schools",
          ["school_id", "school_name", "note", "year_end"],
          [list(s) for s in SCHOOLS],
          [12, 20, 20, 10], "tbl_schools")

add_sheet("classes",
          ["class_id", "class_name", "program", "subject", "teacher_id", "fee_per_session"],
          [list(c) for c in CLASSES],
          [10, 30, 12, 14, 11, 16], "tbl_classes", number_cols=[6])

add_sheet("teachers",
          ["teacher_id", "teacher_name", "note"],
          [list(t) for t in TEACHERS],
          [12, 18, 34], "tbl_teachers")

wb.save(OUT)
print(f"저장 완료: {OUT}")
print(f"  students     : {len(student_rows)} (원생 {len(enrolled)} / 예비원생 {len(prospective)})")
print(f"  consultations: {len(consult_rows)}")
print(f"  monthly_fee 채워짐: {sum(1 for r in student_rows if r[15])}")
print(f"  schools {len(SCHOOLS)} / classes {len(CLASSES)} / teachers {len(TEACHERS)}")
